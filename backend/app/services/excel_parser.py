"""Excel workbook parser with an explicit column-mapping layer.

The mapping is intentionally decoupled from the raw column headers so the
SharePoint workbook's Danish column names can change without touching the
parsing logic. Every row is processed; rows that fail are captured as
RowIssue entries rather than dropped silently.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

from openpyxl import load_workbook

from app.schemas.cases import Case, CaseStatus, GeocodeStatus, RowIssue, RowIssueSeverity
from app.services.address import AddressNormalizationError, normalize_address

# Maps our internal field name -> the set of accepted source column header
# variants (case-insensitive, whitespace-trimmed). Add variants here as the
# real SharePoint sheet's headers become known.
COLUMN_MAPPING: dict[str, list[str]] = {
    "case_number": ["sagsnummer", "sagsnr", "sag nr", "case number"],
    "customer_name": ["kunde", "kundenavn", "navn"],
    "address": ["adresse", "vejnavn og nr", "address"],
    "advisor": ["rådgiver", "sagsbehandler", "advisor"],
    "department": ["afdeling", "department"],
    "status": ["status"],
    "estimated_value_dkk": ["forventet værdi", "værdi", "estimated value", "beløb"],
    "created_date": ["oprettet", "dato", "created"],
    "note": ["note", "bemærkning", "kommentar"],
}

STATUS_MAPPING: dict[str, CaseStatus] = {
    "ny": CaseStatus.NEW,
    "i gang": CaseStatus.IN_PROGRESS,
    "igangværende": CaseStatus.IN_PROGRESS,
    "tilbud sendt": CaseStatus.OFFER_SENT,
    "tilbud afgivet": CaseStatus.OFFER_SENT,
    "vundet": CaseStatus.WON,
    "tabt": CaseStatus.LOST,
    "afventer": CaseStatus.ON_HOLD,
    "på hold": CaseStatus.ON_HOLD,
}


class WorkbookParseError(RuntimeError):
    """Raised when the workbook structure itself is unreadable (wrong sheet, no header, etc.)."""


@dataclass
class ParseResult:
    cases: list[Case] = field(default_factory=list)
    issues: list[RowIssue] = field(default_factory=list)
    total_rows: int = 0


def _normalize_header(value: object) -> str:
    return str(value or "").strip().lower()


def _build_header_index(header_row: tuple) -> dict[str, int]:
    """Returns internal-field-name -> column index, based on COLUMN_MAPPING."""
    normalized_headers = {_normalize_header(cell): idx for idx, cell in enumerate(header_row)}
    index: dict[str, int] = {}
    for field_name, variants in COLUMN_MAPPING.items():
        for variant in variants:
            if variant in normalized_headers:
                index[field_name] = normalized_headers[variant]
                break
    return index


def _cell_str(row: tuple, idx: int | None) -> str | None:
    if idx is None or idx >= len(row):
        return None
    value = row[idx]
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_status(raw: str | None) -> CaseStatus | None:
    if not raw:
        return None
    return STATUS_MAPPING.get(raw.strip().lower())


def _parse_value(raw: str | None) -> float | None:
    if not raw:
        return None
    cleaned = raw.replace("kr.", "").replace("kr", "").replace(".", "").replace(",", ".").strip()
    try:
        return float(cleaned)
    except ValueError:
        return None


def parse_workbook(content: bytes, worksheet_name: str) -> ParseResult:
    """Parses the given .xlsx bytes into Case records plus row-level issues."""
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises various error types for bad files
        raise WorkbookParseError(f"Kunne ikke læse Excel-filen: {exc}") from exc

    if worksheet_name in workbook.sheetnames:
        sheet = workbook[worksheet_name]
    elif workbook.sheetnames:
        sheet = workbook[workbook.sheetnames[0]]
    else:
        raise WorkbookParseError("Regnearket indeholder ingen arkfaner")

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise WorkbookParseError("Arkfanen er tom (ingen overskriftsrække fundet)") from None

    header_index = _build_header_index(header_row)
    if "address" not in header_index:
        raise WorkbookParseError(
            "Kunne ikke finde en adressekolonne. Forventede en af: "
            + ", ".join(COLUMN_MAPPING["address"])
        )

    result = ParseResult()

    for offset, row in enumerate(rows_iter, start=2):  # row 1 is header
        if row is None or all(v is None for v in row):
            continue
        result.total_rows += 1

        address_raw = _cell_str(row, header_index.get("address"))
        if not address_raw:
            result.issues.append(
                RowIssue(
                    row_number=offset,
                    severity=RowIssueSeverity.ERROR,
                    field="address",
                    message="Adressefelt er tomt",
                )
            )
            continue

        case_number = _cell_str(row, header_index.get("case_number"))
        customer_name = _cell_str(row, header_index.get("customer_name"))
        advisor = _cell_str(row, header_index.get("advisor"))
        department = _cell_str(row, header_index.get("department"))
        status_raw = _cell_str(row, header_index.get("status"))
        value_raw = _cell_str(row, header_index.get("estimated_value_dkk"))
        created_date = _cell_str(row, header_index.get("created_date"))
        note = _cell_str(row, header_index.get("note"))

        status = _parse_status(status_raw)
        if status_raw and status is None:
            result.issues.append(
                RowIssue(
                    row_number=offset,
                    severity=RowIssueSeverity.WARNING,
                    field="status",
                    message=f"Ukendt statusværdi '{status_raw}' blev ikke oversat",
                )
            )

        case_id = case_number or f"row-{offset}"

        try:
            normalized = normalize_address(address_raw)
            case = Case(
                id=case_id,
                case_number=case_number,
                customer_name=customer_name,
                address_raw=address_raw,
                address_normalized=normalized.normalized,
                street=normalized.street,
                house_number=normalized.house_number,
                postal_code=normalized.postal_code,
                city=normalized.city,
                advisor=advisor,
                department=department,
                status=status,
                estimated_value_dkk=_parse_value(value_raw),
                created_date=created_date,
                note=note,
                geocode_status=GeocodeStatus.PENDING,
                source_row=offset,
            )
            result.cases.append(case)
        except AddressNormalizationError as exc:
            result.issues.append(
                RowIssue(
                    row_number=offset,
                    severity=RowIssueSeverity.ERROR,
                    field="address",
                    message=str(exc),
                )
            )
            # Still include the case so it shows up as "unresolved" rather than vanishing.
            result.cases.append(
                Case(
                    id=case_id,
                    case_number=case_number,
                    customer_name=customer_name,
                    address_raw=address_raw,
                    advisor=advisor,
                    department=department,
                    status=status,
                    estimated_value_dkk=_parse_value(value_raw),
                    created_date=created_date,
                    note=note,
                    geocode_status=GeocodeStatus.FAILED,
                    source_row=offset,
                )
            )

    return result
